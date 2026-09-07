import base64
from pathlib import Path

import pytest
from docx import Document

from app.services.document_generator import _insert_signature_in_paragraph, _replace_in_paragraph
from app.services.storage import GENERATED_ROOT, safe_stored_path


PNG = base64.b64encode(base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)).decode()


def test_docx_run_replacement_preserves_surrounding_label():
    paragraph = Document().add_paragraph()
    paragraph.add_run("Full Name: ").bold = True
    paragraph.add_run("Jane Doe")
    assert _replace_in_paragraph(paragraph, "Jane Doe", "Ali Khan")
    assert paragraph.text == "Full Name: Ali Khan"
    assert paragraph.runs[0].bold is True


def test_signature_insertion_does_not_clear_label():
    paragraph = Document().add_paragraph("Applicant Signature: ______")
    assert _insert_signature_in_paragraph(paragraph, "______", PNG)
    assert paragraph.text == "Applicant Signature: "
    assert paragraph._p.xpath(".//a:blip")


def test_stored_paths_cannot_escape_root(tmp_path):
    root = tmp_path.resolve()
    allowed = root / "file.pdf"
    assert safe_stored_path(str(allowed), root) == allowed
    with pytest.raises(ValueError, match="outside"):
        safe_stored_path(str(Path(root).parent / "secret"), root)
