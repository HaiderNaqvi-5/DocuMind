from pathlib import Path

from docx import Document


def extract_docx_content(file_path: str | Path) -> dict:
    path = Path(file_path)
    if path.suffix.lower() != ".docx":
        raise ValueError("Expected a DOCX document.")
    try:
        document = Document(path)
    except Exception as exc:
        raise ValueError("Could not read DOCX document.") from exc

    paragraphs = [
        {"location": f"paragraph:{index}", "text": paragraph.text.strip()}
        for index, paragraph in enumerate(document.paragraphs) if paragraph.text.strip()
    ]
    tables = []
    table_text = []
    for table_index, table in enumerate(document.tables):
        rows = []
        for row_index, row in enumerate(table.rows):
            cells = []
            for cell_index, cell in enumerate(row.cells):
                text = cell.text.strip()
                cells.append({"location": f"table:{table_index}/row:{row_index}/cell:{cell_index}", "text": text})
                if text:
                    table_text.append(text)
            rows.append({"cells": cells})
        tables.append({"rows": rows})
    combined = "\n".join([item["text"] for item in paragraphs] + table_text)
    return {"filename": path.name, "structure": {"paragraphs": paragraphs, "tables": tables},
            "combined_text": combined, "paragraph_count": len(paragraphs), "table_count": len(tables)}


def extract_pdf_content(file_path: str | Path) -> dict:
    import pymupdf

    path = Path(file_path)
    if path.suffix.lower() != ".pdf":
        raise ValueError("Expected a PDF document.")
    try:
        with pymupdf.open(path) as document:
            pages = [{"page": index, "text": page.get_text("text").strip()} for index, page in enumerate(document)]
    except Exception as exc:
        raise ValueError("Could not read PDF document.") from exc
    combined = "\n".join(page["text"] for page in pages if page["text"])
    return {"filename": path.name, "structure": {"pages": pages}, "combined_text": combined,
            "page_count": len(pages)}


def extract_document_content(file_path: str | Path) -> dict:
    suffix = Path(file_path).suffix.lower()
    if suffix == ".docx":
        return extract_docx_content(file_path)
    if suffix == ".pdf":
        return extract_pdf_content(file_path)
    raise ValueError("Only DOCX and PDF documents are supported.")
