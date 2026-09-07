import base64

from io import BytesIO
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.shared import Inches

from app.services.storage import GENERATED_ROOT


GENERATED_DIR = GENERATED_ROOT

GENERATED_DIR.mkdir(
    parents=True,
    exist_ok=True,
)


def _replace_in_paragraph(
    paragraph,
    old_text: str,
    new_text: str,
) -> bool:
    """
    Replace text inside one paragraph while preserving
    existing run formatting as much as possible.
    """

    if old_text not in paragraph.text:
        return False

    runs = paragraph.runs

    if not runs:
        return False


    full_text = "".join(
        run.text
        for run in runs
    )


    start_index = full_text.find(
        old_text
    )

    if start_index == -1:
        return False


    end_index = (
        start_index
        + len(old_text)
    )


    current_position = 0

    start_run_index = None
    end_run_index = None

    start_offset = None
    end_offset = None


    for index, run in enumerate(
        runs
    ):
        run_start = (
            current_position
        )

        run_end = (
            current_position
            + len(run.text)
        )


        if (
            start_run_index is None
            and start_index < run_end
        ):
            start_run_index = index

            start_offset = (
                start_index
                - run_start
            )


        if (
            end_index <= run_end
            and end_run_index is None
        ):
            end_run_index = index

            end_offset = (
                end_index
                - run_start
            )

            break


        current_position = (
            run_end
        )


    if (
        start_run_index is None
        or end_run_index is None
    ):
        return False


    if (
        start_run_index
        == end_run_index
    ):
        run = runs[
            start_run_index
        ]

        before = run.text[
            :start_offset
        ]

        after = run.text[
            end_offset:
        ]

        run.text = (
            before
            + new_text
            + after
        )

        return True


    first_run = runs[
        start_run_index
    ]

    last_run = runs[
        end_run_index
    ]


    before = first_run.text[
        :start_offset
    ]

    after = last_run.text[
        end_offset:
    ]


    first_run.text = (
        before
        + new_text
    )


    for index in range(
        start_run_index + 1,
        end_run_index,
    ):
        runs[index].text = ""


    last_run.text = (
        after
    )


    return True


def _set_paragraph_text(
    paragraph,
    value: str,
) -> None:
    """
    Replace the complete contents of a paragraph.

    The first existing run is reused when possible
    so basic formatting survives.
    """

    runs = paragraph.runs


    if runs:
        runs[0].text = value

        for run in runs[1:]:
            run.text = ""

        return


    paragraph.add_run(
        value
    )


def _normalized_lines(
    value: str,
) -> list[str]:
    return [
        line.strip()
        for line in value.splitlines()
        if line.strip()
    ]


def _replace_block_in_paragraphs(
    paragraphs,
    old_text: str,
    new_text: str,
) -> bool:
    """
    Replace a source block spanning multiple consecutive
    paragraphs.

    The parser joins paragraphs with newline characters,
    so a textarea/block field can point to several existing
    paragraphs.

    The replacement is placed into the first matched
    paragraph. Remaining source paragraphs are cleared.
    """

    old_lines = _normalized_lines(
        old_text
    )


    if len(old_lines) < 2:
        return False


    visible = []

    for paragraph in paragraphs:
        text = (
            paragraph.text
            .strip()
        )

        if not text:
            continue

        visible.append(
            (
                paragraph,
                text,
            )
        )


    required_count = len(
        old_lines
    )


    if (
        len(visible)
        < required_count
    ):
        return False


    for start in range(
        0,
        len(visible)
        - required_count
        + 1,
    ):
        candidate = [
            visible[
                start + offset
            ][1]
            for offset in range(
                required_count
            )
        ]


        if candidate != old_lines:
            continue


        matched_paragraphs = [
            visible[
                start + offset
            ][0]
            for offset in range(
                required_count
            )
        ]


        _set_paragraph_text(
            matched_paragraphs[0],
            new_text,
        )


        for paragraph in (
            matched_paragraphs[1:]
        ):
            _set_paragraph_text(
                paragraph,
                "",
            )


        return True


    return False


def _decode_signature(
    signature_data: str,
) -> BytesIO:
    if not signature_data:
        raise ValueError(
            "Signature data is empty."
        )


    if "," in signature_data:
        header, encoded = (
            signature_data.split(
                ",",
                1,
            )
        )

        if (
            "base64"
            not in header.lower()
        ):
            raise ValueError(
                "Invalid signature data."
            )

    else:
        encoded = (
            signature_data
        )


    try:
        image_bytes = (
            base64.b64decode(
                encoded,
                validate=True,
            )
        )

    except Exception as exc:
        raise ValueError(
            "Invalid base64 signature."
        ) from exc


    if not image_bytes.startswith(
        b"\x89PNG\r\n\x1a\n"
    ):
        raise ValueError(
            "Signature must be a PNG image."
        )

    if len(image_bytes) > 2 * 1024 * 1024:
        raise ValueError("Signature image is too large.")


    image_stream = BytesIO(
        image_bytes
    )

    image_stream.seek(
        0
    )

    return image_stream


def _insert_signature_in_paragraph(
    paragraph,
    placeholder: str,
    signature_data: str,
) -> bool:
    if placeholder not in paragraph.text:
        return False


    signature_stream = (
        _decode_signature(
            signature_data
        )
    )


    # Remove only the detected marker; labels and surrounding runs remain intact.
    if not _replace_in_paragraph(paragraph, placeholder, ""):
        return False
    run = paragraph.add_run()


    run.add_picture(
        signature_stream,
        width=Inches(1.7),
    )


    return True


def _process_paragraph_collection(
    paragraphs,
    text_replacements: dict[str, str],
    signature_replacements: dict[str, str],
) -> tuple[int, int]:
    text_count = 0
    signature_count = 0


    # ==================================================
    # BLOCK REPLACEMENTS FIRST
    # ==================================================

    block_replacements = {
        old_text: new_text
        for old_text, new_text
        in text_replacements.items()
        if "\n" in old_text
    }


    for (
        old_text,
        new_text,
    ) in block_replacements.items():
        if _replace_block_in_paragraphs(
            paragraphs,
            old_text,
            new_text,
        ):
            text_count += 1


    # ==================================================
    # NORMAL SINGLE-PARAGRAPH REPLACEMENTS
    # ==================================================

    scalar_replacements = {
        old_text: new_text
        for old_text, new_text
        in text_replacements.items()
        if "\n" not in old_text
    }


    for paragraph in paragraphs:
        for (
            old_text,
            new_text,
        ) in scalar_replacements.items():
            if _replace_in_paragraph(
                paragraph,
                old_text,
                new_text,
            ):
                text_count += 1


        for (
            placeholder,
            signature_data,
        ) in (
            signature_replacements
            .items()
        ):
            if _insert_signature_in_paragraph(
                paragraph,
                placeholder,
                signature_data,
            ):
                signature_count += 1


    return (
        text_count,
        signature_count,
    )


def _process_table(
    table,
    text_replacements:
        dict[str, str],
    signature_replacements:
        dict[str, str],
) -> tuple[int, int]:
    text_count = 0
    signature_count = 0


    for row in table.rows:
        for cell in row.cells:
            (
                cell_text_count,
                cell_signature_count,
            ) = _process_paragraph_collection(
                cell.paragraphs,
                text_replacements,
                signature_replacements,
            )


            text_count += (
                cell_text_count
            )

            signature_count += (
                cell_signature_count
            )


            for nested_table in (
                cell.tables
            ):
                (
                    nested_text,
                    nested_signature,
                ) = _process_table(
                    nested_table,
                    text_replacements,
                    signature_replacements,
                )


                text_count += (
                    nested_text
                )

                signature_count += (
                    nested_signature
                )


    return (
        text_count,
        signature_count,
    )


def generate_docx_document(
    template_path:
        str | Path,
    replacements:
        dict[str, str],
    signature_replacements:
        dict[str, str] | None = None,
    output_name:
        str | None = None,
) -> dict:
    template_path = Path(
        template_path
    )


    if not template_path.exists():
        raise FileNotFoundError(
            f"Template not found: {template_path}"
        )


    if (
        template_path
        .suffix
        .lower()
        != ".docx"
    ):
        raise ValueError(
            "Only DOCX templates are supported."
        )


    document = Document(
        str(
            template_path
        )
    )


    signature_replacements = (
        signature_replacements
        or {}
    )


    text_count = 0
    signature_count = 0


    # ==================================================
    # BODY
    # ==================================================

    (
        body_text_count,
        body_signature_count,
    ) = _process_paragraph_collection(
        document.paragraphs,
        replacements,
        signature_replacements,
    )


    text_count += (
        body_text_count
    )

    signature_count += (
        body_signature_count
    )


    # ==================================================
    # TABLES
    # ==================================================

    for table in document.tables:
        (
            table_text_count,
            table_signature_count,
        ) = _process_table(
            table,
            replacements,
            signature_replacements,
        )


        text_count += (
            table_text_count
        )

        signature_count += (
            table_signature_count
        )


    # ==================================================
    # SAVE
    # ==================================================

    if output_name:
        safe_name = (
            Path(
                output_name
            )
            .stem
        )

        filename = (
            f"{uuid4().hex}_"
            f"{safe_name}.docx"
        )

    else:
        filename = (
            f"{uuid4().hex}_"
            "generated.docx"
        )


    output_path = (
        GENERATED_DIR
        / filename
    )


    document.save(
        str(
            output_path
        )
    )


    return {
        "path":
            str(
                output_path
            ),

        "filename":
            filename,

        "replacement_count":
            text_count,

        "signature_count":
            signature_count,
    }


# Compatibility for existing callers; new code dispatches through document_service.
generate_document = generate_docx_document
